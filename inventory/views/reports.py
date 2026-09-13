from django.contrib.auth.decorators import login_required, permission_required
from django.shortcuts import render
from django.http import HttpResponse
from django.db.models import Sum, F, Q
import csv

from ..models import Product, Category, Warehouse

@login_required
@permission_required("inventory.view_inventorytransaction", raise_exception=True)
def inventory_report(request):

    search = request.GET.get("search", "")
    category = request.GET.get("category", "")
    status = request.GET.get("status", "")
    warehouse = request.GET.get("warehouse", "")

    products = Product.objects.select_related(
        "category",
        "color",
        "size",
    )
    
    if warehouse:
        products = products.filter(stocks__warehouse_id=warehouse).distinct()

    if search:
        products = products.filter(
            Q(name__icontains=search) |
            Q(id__icontains=search)
        )

    if category:
        products = products.filter(category_id=category)

    if status == "low":
        products = products.filter(quantity__lte=F("minimum_stock"), quantity__gt=0)
    elif status == "out":
        products = products.filter(quantity=0)

    products = products.annotate(total_value=F("cost_price") * F("quantity")).order_by("-id")
    products_list = list(products)

    total_quantity = sum(p.quantity for p in products_list)
    inventory_value = sum(p.total_value for p in products_list if p.total_value)

    if request.GET.get("export") == "xlsx":
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Inventory Report"
        ws.sheet_view.rightToLeft = True

        header_fill = PatternFill(start_color="1A2738", end_color="1A2738", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True, size=12)
        center_align = Alignment(horizontal="center", vertical="center")
        left_align = Alignment(horizontal="left", vertical="center")
        thin_border = Border(left=Side(style='thin', color='E2E4E0'), 
                             right=Side(style='thin', color='E2E4E0'), 
                             top=Side(style='thin', color='E2E4E0'), 
                             bottom=Side(style='thin', color='E2E4E0'))

        headers = ['كود المنتج', 'الاسم', 'التصنيف', 'السعر', 'الكمية', 'الحد الأدنى', 'الحالة', 'القيمة الإجمالية']
        
        ws.append(headers)
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align
            cell.border = thin_border
            ws.column_dimensions[get_column_letter(col_num)].width = 20

        ws.column_dimensions['B'].width = 35

        for row_num, p in enumerate(products_list, 2):
            row_data = [
                p.id,
                p.name,
                p.category.name if p.category else 'بدون تصنيف',
                float(p.cost_price),
                p.quantity,
                p.minimum_stock,
                p.stock_status_text,
                float(p.total_value) if p.total_value else 0.0
            ]
            
            for col_num, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_num, column=col_num, value=value)
                cell.border = thin_border
                if col_num in [4, 8]:
                    cell.number_format = '#,##0.00'
                    cell.alignment = center_align
                elif col_num in [5, 6]:
                    cell.alignment = center_align
                else:
                    cell.alignment = left_align

                if col_num == 7:
                    cell.alignment = center_align
                    if p.quantity == 0:
                        cell.font = Font(color="EF4444", bold=True)
                    elif p.quantity <= p.minimum_stock:
                        cell.font = Font(color="F59E0B", bold=True)
                    else:
                        cell.font = Font(color="10B981", bold=True)

        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="inventory_report.xlsx"'
        wb.save(response)
        return response

    context = {
        "products": products_list,
        "categories": Category.objects.all(),
        "warehouses": Warehouse.objects.all(),
        "search": search,
        "category": category,
        "status": status,
        "warehouse": warehouse,
        "total_products": len(products_list),
        "total_quantity": total_quantity,
        "inventory_value": inventory_value,
        "low_stock": Product.objects.filter(quantity__lte=F("minimum_stock"), quantity__gt=0).count(),
        "out_of_stock": Product.objects.filter(quantity=0).count(),
    }
    return render(request, "reports/inventory.html", context)
